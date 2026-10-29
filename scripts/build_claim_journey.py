"""Extract IRDAI Handbook Table 66 into a transparent, source-preserving JSON dataset.

The source contains annual aggregate claim development and aging for health insurance
(excluding personal accident and travel). Values are preserved as published. We do not
force the flow to reconcile; the site displays a warning when a simple reconciliation
check does not balance.
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

YEARS_RE = re.compile(r"^20\d\d-\d\d$")
FLOW_ROWS = {
    8: "opening",
    9: "new_registered",
    10: "paid",
    11: "disallowed",
    12: "repudiated",
    13: "closing",
}
AGING_ROWS = {
    "paid": [(r, _label) for r, _label in [
        (16, "Within 1 month"), (17, "1–3 months"), (18, "3–6 months"),
        (19, "6–12 months"), (20, "1–2 years"), (21, "Beyond 2 years")]],
    "repudiated": [(r, _label) for r, _label in [
        (24, "Within 1 month"), (25, "1–3 months"), (26, "3–6 months"),
        (27, "6–12 months"), (28, "1–2 years"), (29, "Beyond 2 years")]],
    "outstanding": [(r, _label) for r, _label in [
        (32, "Within 1 month"), (33, "1–3 months"), (34, "3–6 months"),
        (35, "6–12 months"), (36, "1–2 years"), (37, "Beyond 2 years")]],
}


def num(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, str) and v.strip() in {"-", "--", "NA", "N.A."}:
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def normalise_category(s):
    s = " ".join(str(s or "").lower().split())
    if s in {"cashless", "only cashless"}:
        return "cashless"
    if s in {"reimbursement", "only reimbursement"}:
        return "reimbursement"
    if s == "both cashless and reimbursement":
        return "both_cashless_and_reimbursement"
    if s == "benefit based":
        return "benefit_based"
    if s == "total":
        return "total"
    return None


def val_pair(rows, row, col):
    """Read a one-based Excel row/column from the selected sheet's value matrix."""
    values = rows[row - 1]
    return {"count": num(values[col - 1] if col - 1 < len(values) else None),
            "amount_lakh": num(values[col] if col < len(values) else None)}


def route_layout(ws, year, start):
    if year == "2015-16":
        routes = {"tpa": (start, 8), "direct": (start + 8, 8), "all": (start + 16, 2)}
    else:
        routes = {"tpa": (start, 10), "direct": (start + 10, 10), "all": (start + 20, 2)}
    return routes


def categories_for_route(rows, year, start, width):
    if width == 2:
        return {"total": start}
    cats = {}
    header = rows[4]  # Excel row 5 contains the categories.
    for off in range(0, width, 2):
        idx = start + off - 1
        key = normalise_category(header[idx] if idx < len(header) else None)
        if key:
            cats[key] = start + off
    return cats


def build(src):
    wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
    if "66" not in wb.sheetnames:
        wb.close()
        raise ValueError("Workbook is missing Table 66 sheet")
    ws = wb["66"]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    years = []
    header_years = rows[2] if len(rows) >= 3 else []
    for col in range(2, len(header_years) + 1):
        value = header_years[col - 1]
        if isinstance(value, str) and YEARS_RE.match(value.strip()):
            years.append((value.strip(), col))
    if len(years) != 10:
        raise ValueError(f"Expected ten annual blocks in Table 66, found {len(years)}: {years}")

    result = {
        "meta": {
            "source": "IRDAI Handbook on Indian Insurance Statistics 2024-25",
            "source_table": "Table 66",
            "title": "Details of Claims Development and Aging — Health Insurance Business (excluding Travel and Personal Accident)",
            "latest_year": years[-1][0],
            "years": [y for y, _ in years],
            "amount_unit": "₹ lakh",
            "count_unit": "claims as reported by IRDAI",
            "aging_notes": {
                "paid": "Paid-claim aging is reckoned from the date of first intimation.",
                "repudiated": "Repudiated-claim aging is reckoned from the date of receipt of the last requirement.",
                "outstanding": "Outstanding-claim aging is reckoned from the date of receipt of the last requirement."
            },
            "limitations": [
                "Annual flows are not a single claim cohort: claims paid or repudiated in a year may have been registered in earlier years.",
                "Table 66 is industry-level data, not insurer-level data.",
                "Count and amount are separate measures and must not be mixed.",
                "The source may not fully reconcile under a simple opening + new = paid + repudiated + closing equation. Values are preserved as reported; no balancing adjustments are made."
            ]
        },
        "years": {}
    }

    labels = {
        8: "opening", 9: "new_registered", 10: "paid", 11: "disallowed",
        12: "repudiated", 13: "closing"
    }
    for idx, (year, start) in enumerate(years):
        width = (years[idx + 1][1] - start) if idx + 1 < len(years) else (18 if year == "2015-16" else 22)
        routes = route_layout(ws, year, start)
        yobj = {"year": year, "routes": {}}
        for route, (rstart, rwidth) in routes.items():
            cat_cols = categories_for_route(rows, year, rstart, rwidth)
            robj = {"categories": {}, "flows": {}, "aging": {}}
            for cat, cstart in cat_cols.items():
                cobj = {"flows": {}, "aging": {}}
                for row, key in labels.items():
                    cobj["flows"][key] = val_pair(rows, row, cstart)
                for ageing_type, buckets in AGING_ROWS.items():
                    cobj["aging"][ageing_type] = [
                        {"label": label, **val_pair(rows, row, cstart)} for row, label in buckets
                    ]
                robj["categories"][cat] = cobj
            robj["flows"] = robj["categories"].get("total", {}).get("flows", {})
            robj["aging"] = robj["categories"].get("total", {}).get("aging", {})
            yobj["routes"][route] = robj
        # Diagnostics only: preserve source values and never rebalance the flow.
        def residual(route, unit):
            f = yobj["routes"][route]["flows"]
            def v(k):
                return (f.get(k) or {}).get(unit)
            vals = [v("opening"), v("new_registered"), v("paid"), v("repudiated"), v("closing")]
            if any(x is None for x in vals):
                return None
            return vals[0] + vals[1] - vals[2] - vals[3] - vals[4]
        yobj["reconciliation"] = {}
        for route in ("all", "tpa", "direct"):
            flow = yobj["routes"][route]["flows"]
            base_residual = residual(route, "amount_lakh")
            disallowed = (flow.get("disallowed") or {}).get("amount_lakh")
            yobj["reconciliation"][route] = {
                "count_residual": residual(route, "count"),
                "amount_lakh_residual": base_residual,
                "amount_lakh_residual_including_disallowed": (base_residual - disallowed) if base_residual is not None and disallowed is not None else None
            }
        result["years"][year] = yobj

    return result


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: python scripts/build_claim_journey.py PartIII.xlsx output.json")
    src, dst = Path(sys.argv[1]), Path(sys.argv[2])
    data = build(src)
    dst.parent.mkdir(parents=True, exist_ok=True)
    dst.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")))
    print(f"Wrote {dst} ({dst.stat().st_size / 1024:.0f} KB), {len(data['years'])} years")
