"""
Build app_data.json for the website from the IRDAI Handbook Excel files.

Usage:
    python scripts/build_app_data.py "<folder with Part II.xlsx and Part III.xlsx>" data/processed/state_health.json docs/data/app_data.json

Sources (IRDAI Handbook on Indian Insurance Statistics 2024-25):
  Table 67 (via state_health.json): state-wise health insurance business
  Table 72: state-wise claim settlement (cashless / reimbursement)
  Table 53: status of claims, insurer-wise (only stand-alone health insurers are used,
            because for multi-line insurers this table mixes motor, fire etc. with health)
  Table 56: status of grievances, insurer-wise
  Table 58: health insurance policies, insurer-wise (denominator for complaints)

Every check prints PASS/FAIL. Any FAIL stops the build (exit 1) and nothing is written.
"""
import json
import re
import sys
from datetime import date
from pathlib import Path

import openpyxl
import pandas as pd

FAILS = []


def check(name, ok, detail=""):
    print(f"[{'PASS' if ok else 'FAIL'}] {name}: {detail}")
    if not ok:
        FAILS.append(name)


def clean(x):
    return " ".join(str(x).replace("\n", " ").split()).strip() if x is not None else ""


def num(x):
    if x is None or (isinstance(x, str) and x.strip() in ("", "-", "--")):
        return None
    try:
        return float(x)
    except (TypeError, ValueError):
        return None


def key_name(n):
    """Normalise insurer names so the same company matches across tables."""
    n = clean(n).lower().replace("&", "and")
    n = re.sub(r"[^a-z ]", " ", n)
    n = re.sub(r"\b(co|company|ltd|limited|insurance|india|the)\b", " ", n)
    return " ".join(n.split())


STATE_FIX = {"orissa": "Odisha", "uttrakhand": "Uttarakhand", "jammu kashmir": "Jammu & Kashmir",
             "jammu & kashmir": "Jammu & Kashmir", "total": "All India"}


def state_name(s):
    s = clean(s)
    return STATE_FIX.get(s.lower(), s)


# ---------------- Table 67 (already validated by process_state_health.py) ----------------
def states_block(state_json):
    rows = json.loads(Path(state_json).read_text())
    df = pd.DataFrame(rows)
    groups = {"individual": "individual", "group": "group", "government_scheme": "government", "total_row": "total"}
    df["g"] = df["group"].map(groups)
    metrics = ["policies", "persons_covered_thousand", "gross_premium_lakh", "claims_paid_count", "claims_paid_lakh"]
    agg = df.groupby(["state_name", "fiscal_year", "g"])[metrics].sum(min_count=1).reset_index()
    # All India = sum of states. process_state_health.py already verified this equals IRDAI's own
    # All India row (275/275 within 1%), so the sum and the official total are interchangeable.
    nat = df.groupby(["fiscal_year", "g"])[metrics].sum(min_count=1).reset_index()
    nat["state_name"] = "All India"
    agg = pd.concat([agg, nat], ignore_index=True)
    # ChatGPT review (Round 6): if any component has a negative premium or claims value, summing it with
    # positive components gives a misleading ratio. Flag those groups so the app shows "Not available".
    neg = df[(df["gross_premium_lakh"] < 0) | (df["claims_paid_lakh"] < 0)]
    neg_keys = set(zip(neg.state_name, neg.fiscal_year, neg.g))
    neg_keys |= {("All India", y, g) for _, y, g in neg_keys}
    out = {}
    for _, r in agg.iterrows():
        d = {m: (None if pd.isna(r[m]) else round(float(r[m]), 3)) for m in metrics}
        d["has_negative_component"] = (r.state_name, r.fiscal_year, r.g) in neg_keys
        out.setdefault(r.state_name, {}).setdefault(r.fiscal_year, {})[r.g] = d
    check("table67_states_loaded", len(out) >= 37 and "All India" in out, f"{len(out)} entries incl. All India")
    print(f"[INFO] groups with a negative component (ratio shown as Not available): {len(neg_keys)}")
    return out


# ---------------- Table 72: cashless vs reimbursement ----------------
def cashless_block(part3):
    rows = list(openpyxl.load_workbook(part3, data_only=True, read_only=True)["72"].iter_rows(values_only=True))
    h = next(i for i, r in enumerate(rows) if clean(r[0]).lower().startswith("s.no"))
    y = biz = typ = mode = None
    cols = []
    for j in range(2, len(rows[h])):
        if rows[h][j] is not None:
            y, biz, typ, mode = clean(rows[h][j]), None, None, None
        if rows[h + 1][j] is not None:
            biz, typ, mode = clean(rows[h + 1][j]), None, None
        if rows[h + 2][j] is not None:
            typ, mode = clean(rows[h + 2][j]), None
        if rows[h + 3][j] is not None:
            mode = clean(rows[h + 3][j])
        met = clean(rows[h + 4][j]).lower()
        if not met.startswith("no. of claims"):
            continue  # use claim counts only
        b = "individual" if biz and biz.lower().startswith("individual") else "group" if biz else None
        if typ == "Total":
            k = "total"
        elif typ == "Indemnity" and mode == "Cashless":
            k = "cashless"
        elif typ == "Indemnity" and mode == "Reimbursement":
            k = "reimbursement"
        elif typ == "Indemnity":
            k = "both"
        elif typ == "Benefit Based":
            k = "benefit"
        else:
            k = None
        if b and k:
            cols.append((j, y, b, k))
    out, bad = {}, []
    for r in rows[h + 5:]:
        if not r[1]:
            continue
        st = state_name(r[1])
        for j, y, b, k in cols:
            out.setdefault(st, {}).setdefault(y, {}).setdefault(b, {})[k] = num(r[j])
    for st, ys in out.items():
        for y, bs in ys.items():
            for b, d in bs.items():
                parts = [d.get(k) for k in ("cashless", "reimbursement", "both", "benefit")]
                if d.get("total") is not None and all(p is not None for p in parts):
                    if abs(sum(parts) - d["total"]) > max(2, 0.005 * d["total"]):
                        bad.append((st, y, b, sum(parts), d["total"]))
                c = d.get("cashless")
                t = sum(parts) if all(p is not None for p in parts) else None
                d["parts_total"] = t
                d["total_mismatch"] = bool(d.get("total") is not None and t is not None and abs(t - d["total"]) > max(2, 0.005 * d["total"]))
                d["cashless_share_pct"] = round(c / t * 100, 1) if t and c is not None and t > 0 and c >= 0 else None
    years = sorted({y for ys in out.values() for y in ys})
    n_all = sum(1 for ys in out.values() for bs in ys.values() for _ in bs)
    by_year = {}
    for st, y, b, *_ in bad:
        by_year[y] = by_year.get(y, 0) + 1
    print(f"[WARN] table72_parts_vs_irdai_total: {len(bad)}/{n_all} rows where IRDAI's Total column differs from its own parts "
          f"(by year {by_year}). Cashless share is computed from the parts and these rows are flagged in the app.")
    check("table72_coverage", len(out) >= 36 and len(years) >= 5, f"{len(out)} rows, years {years}")
    return out


# ---------------- Tables 53 + 56 + 58: stand-alone health insurers ----------------
def section_rows(ws_rows, name_col, start_label="stand-alone health insurers"):
    """Yield rows between the 'Stand-alone Health Insurers' heading and its Total row."""
    inside = False
    for r in ws_rows:
        n = clean(r[name_col]).lower() or (clean(r[name_col - 1]).lower() if name_col > 0 else "")
        if n.startswith(start_label) and "total" not in n:
            inside = True
            continue
        if inside and n.startswith(start_label) and "total" in n:
            yield ("TOTAL", r)
            return
        if inside and n:
            yield (clean(r[name_col]), r)


def find_year_cols(rows, year, header_label_row):
    hr = next(i for i, r in enumerate(rows) if clean(r[0]).lower().startswith("s.no"))
    y = None
    cols = {}
    for j, x in enumerate(rows[hr]):
        if x is not None:
            y = clean(x)
        if y == year:
            cols[j] = clean(rows[hr + header_label_row][j]).lower()
    return hr, cols


def insurers_block(part2, part3, year="2024-25"):
    wb2 = openpyxl.load_workbook(part2, data_only=True, read_only=True)
    wb3 = openpyxl.load_workbook(part3, data_only=True, read_only=True)

    # Table 53
    r53 = list(wb2["53"].iter_rows(values_only=True))
    _, c53 = find_year_cols(r53, year, 2)
    want = {"claims o/s at start": "os_start", "claims intimated": "intimated", "claims paid": "paid",
            "claims repudiated": "repudiated", "claims closed": "closed", "claims o/s at the end": "os_end"}
    colmap = {}
    for j, lab in c53.items():
        for k, v in want.items():
            if lab.startswith(k) and v not in colmap:
                colmap[v] = j
    check("table53_columns_found", len(colmap) == 6, f"{sorted(colmap)}")
    ins, total53 = {}, None
    for n, r in section_rows(r53, 1):
        d = {k: num(r[j]) for k, j in colmap.items()}
        if n == "TOTAL":
            total53 = d
        else:
            ins[key_name(n)] = {"name": re.sub(r"[#$~^*@%]+$", "", n).strip(), **d}
    sums = {k: sum(v[k] or 0 for v in ins.values()) for k in colmap}
    ok = total53 and all(abs(sums[k] - (total53[k] or 0)) <= max(2, 0.001 * (total53[k] or 0)) for k in colmap)
    check("table53_insurers_add_up_to_total", bool(ok), f"{len(ins)} insurers")
    # Flow check: start + intimated = paid + repudiated + closed + end (IRDAI note says O/S may differ slightly)
    flow_bad = []
    for v in ins.values():
        if None in (v["os_start"], v["intimated"], v["paid"], v["repudiated"], v["closed"], v["os_end"]):
            continue
        lhs, rhs = v["os_start"] + v["intimated"], v["paid"] + v["repudiated"] + v["closed"] + v["os_end"]
        if abs(lhs - rhs) > 0.02 * max(lhs, 1):
            flow_bad.append(v["name"])
    print(f"[INFO] table53_claim_flow: {len(ins)-len(flow_bad)}/{len(ins)} insurers balance within 2% "
          f"(IRDAI note: O/S at end may differ due to reopened claims). Off: {flow_bad}")

    # Table 56 (2024-25 is a separate block with its own name column)
    r56 = list(wb2[[s for s in wb2.sheetnames if s.strip() == "56"][0]].iter_rows(values_only=True))
    hr = next(i for i, r in enumerate(r56) if clean(r[0]).lower().startswith("s.no"))
    yr_start = next(j for j, x in enumerate(r56[hr]) if clean(x) == year)
    name_col = next(j for j in range(yr_start, len(r56[hr + 1])) if clean(r56[hr + 1][j]).lower() == "insurer")
    rep_col = next(j for j in range(name_col, len(r56[hr + 1])) if clean(r56[hr + 1][j]).lower().startswith("reported"))
    open_col = next(j for j in range(name_col, len(r56[hr + 1])) if clean(r56[hr + 1][j]).lower().startswith("opening"))
    # IRDAI footnote: "Reported during the year" includes the opening balance, so new complaints = reported - opening
    foot = " ".join(clean(x) for r in r56 for x in r if x and "including opening balance" in clean(x).lower())
    check("table56_footnote_understood", "including opening balance" in foot.lower(), "reported includes opening balance")
    att_col = next(j for j in range(name_col, len(r56[hr + 1])) if clean(r56[hr + 1][j]).lower().startswith("attended"))
    comp = {}
    for n, r in section_rows(r56, name_col):
        if n != "TOTAL":
            rep, opn = num(r[rep_col]), num(r[open_col])
            new = rep - (opn or 0) if rep is not None else None
            comp[key_name(n)] = (new, num(r[att_col]))
    check("table56_complaints_found", len(comp) >= 6, f"{len(comp)} insurers, col {rep_col}")

    # Table 58 (policies, TOTAL block of the year)
    r58 = list(wb3["58"].iter_rows(values_only=True))
    hr58 = next(i for i, r in enumerate(r58) if clean(r[0]).lower().startswith("s.no"))
    ycol = next(j for j, x in enumerate(r58[hr58]) if clean(x) == year)
    tot_col = next(j for j in range(ycol, len(r58[hr58 + 1])) if clean(r58[hr58 + 1][j]).upper() == "TOTAL")
    assert clean(r58[hr58 + 2][tot_col]).lower().startswith("no. of policies")
    pol = {}
    for n, r in section_rows(r58, 1):
        if n != "TOTAL":
            pol[key_name(n)] = num(r[tot_col])
    check("table58_policies_found", len(pol) >= 6, f"{len(pol)} insurers")

    out = []
    unmatched = []
    for k, v in ins.items():
        decided = sum(x or 0 for x in (v["paid"], v["repudiated"], v["closed"]))
        handled = (v["os_start"] or 0) + (v["intimated"] or 0)
        v["repudiated_pct"] = round(v["repudiated"] / decided * 100, 1) if decided > 0 and v["repudiated"] is not None else None
        v["pending_end_pct"] = round(v["os_end"] / handled * 100, 1) if handled > 0 and v["os_end"] is not None else None
        c = comp.get(k)
        p = pol.get(k)
        if c is None or p is None:
            unmatched.append(v["name"])
        v["complaints_new"] = c[0] if c else None
        v["policies"] = p
        v["complaints_per_10k"] = round(c[0] / p * 10000, 1) if c and c[0] is not None and p and p > 0 else None
        v["new_company"] = any(w in v["name"].lower() for w in ("galaxy", "narayana"))  # started 2024; tiny base
        v["claims_handled"] = handled
        out.append(v)
    check("insurer_names_match_across_tables", len(unmatched) <= 2, f"unmatched: {unmatched}")
    return sorted(out, key=lambda v: -(v["claims_handled"] or 0))


if __name__ == "__main__":
    folder, state_json, dst = Path(sys.argv[1]), sys.argv[2], Path(sys.argv[3])
    data = {
        "meta": {
            "source": "IRDAI Handbook on Indian Insurance Statistics 2024-25",
            "source_url": "https://irdai.gov.in/handbook-of-indian-insurance",
            "tables": {"states": "Table 67", "cashless": "Table 72", "insurers": "Tables 53, 56, 58"},
            "latest_year": "2024-25",
            "built_on": date.today().isoformat(),
        },
        "states": states_block(state_json),
        "cashless": cashless_block(folder / "Part III.xlsx"),
        "insurers": insurers_block(folder / "Part II.xlsx", folder / "Part III.xlsx"),
    }
    # Cross-table check: All India individual claims count in Table 72 must equal Table 67
    try:
        t72 = data["cashless"]["All India"]["2024-25"]["individual"]["parts_total"]
        t67 = data["states"]["All India"]["2024-25"]["individual"]["claims_paid_count"]
        check("table72_matches_table67_all_india", abs(t72 - t67) <= 1, f"Table 72: {t72:,.0f} vs Table 67: {t67:,.0f}")
    except KeyError as e:
        check("table72_matches_table67_all_india", False, f"missing {e}")
    if FAILS:
        sys.exit(f"\nBUILD STOPPED. Failed checks: {FAILS}. Nothing was written.")
    dst.parent.mkdir(parents=True, exist_ok=True)
    dst.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")))
    print(f"\nWrote {dst} ({dst.stat().st_size/1024:.0f} KB)")
